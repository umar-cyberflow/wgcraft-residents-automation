# -*- coding: utf-8 -*-
"""
Selenium automation for a billing web app (WG Craft): restore the number of
residents ("Проживает") on subscriber accounts after a bulk change.

For every account in an Excel list:
  1) Subscribers -> search by account number
  2) Create a registration document (number = 1, date = today) and confirm
  3) Open the residents dialog, set the new count (3/4/5 from a repeating
     pattern), fill the end date with today if it is empty, confirm

Usage:
  python set_residents.py --dry-run   # no browser: show accounts and planned values
  python set_residents.py             # TEST: only the first 3 accounts
  python set_residents.py --all       # everything (resumes from progress.csv)
"""
import argparse
import csv
import datetime as dt
import sys
import time
from pathlib import Path

from dotenv import dotenv_values
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import (
    ElementClickInterceptedException,
    ElementNotInteractableException,
    StaleElementReferenceException,
    TimeoutException,
)
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

# =====================  SETTINGS  =====================
BASE_DIR = Path(__file__).resolve().parent

CFG = dotenv_values(BASE_DIR / ".env")
BASE_URL = (CFG.get("WG_BASE_URL") or "https://wgcraft.example.com").rstrip("/")
SEARCH_URL = BASE_URL + "/fcmf"
BROWSER = "chrome"                    # "chrome" or "edge"

EXCEL_PATH = BASE_DIR / "accounts.xlsx"   # filtered workbook (or pass --excel)
SHEET_NAME = None                     # None = active sheet
ACCOUNT_COL = "K"                     # column with the account number ("Лиц.счет")
FIRST_DATA_ROW = 6                    # first data row (below header and filter rows)

# (value, how many accounts in a row) -- the pattern repeats when it ends
PATTERN = [
    (4, 10), (3, 1), (4, 5), (5, 2),
    (4, 10), (3, 1), (4, 5), (5, 1),
]

REG_NUMBER = "1"                      # value typed into the registration-number field
EXPECTED_OLD = "1"                    # only change accounts whose current value is this; others are skipped
                                      # (None = no check)

T_WAIT = 40                           # max wait per step in seconds (the site can be slow)
PAUSE = 0.4                           # short pause between clicks
MAX_CONSECUTIVE_ERRORS = 5            # stop after this many errors in a row

CRED_FILE = BASE_DIR / ".env"
PROGRESS_FILE = BASE_DIR / "progress.csv"
ERRORS_DIR = BASE_DIR / "errors"

# Statuses that are NOT retried when the script is restarted
FINAL_STATUSES = {"done", "unverified", "skipped_not_1", "error_partial"}

MODAL_REG_TITLE = "Рег №"
MODAL_RES_TITLE = "Изменить кол. проживающих"
# ========================================================

driver = None


# ---------------------  Excel and value pattern  ---------------------
def norm_account(v):
    """Return the account number as text; restore a lost leading zero."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip().replace(" ", "")
    if not s.isdigit():
        return None
    if len(s) < 10:
        s = s.zfill(10)
    return s


def read_accounts(path):
    """Read account numbers from VISIBLE (filtered-in) rows only."""
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.active
    dims = ws.row_dimensions
    accounts, seen = [], set()
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if r in dims and dims[r].hidden:
            continue
        acc = norm_account(ws[f"{ACCOUNT_COL}{r}"].value)
        if acc and acc not in seen:
            seen.add(acc)
            accounts.append(acc)
    return accounts


EXPANDED = [val for val, cnt in PATTERN for _ in range(cnt)]


def value_for(position):
    """position = index of the account in the list (from 0); the pattern repeats."""
    return EXPANDED[position % len(EXPANDED)]


# ---------------------  Progress file  ---------------------
def load_progress():
    status = {}
    if PROGRESS_FILE.exists():
        with open(PROGRESS_FILE, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                status[row["account"]] = row["status"]
    return status


def log_progress(acc, pos, st, val, msg=""):
    new = not PROGRESS_FILE.exists()
    with open(PROGRESS_FILE, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["account", "position", "status", "value", "message", "time"])
        w.writerow([acc, pos, st, val, msg, dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])


# ---------------------  Selenium helpers  ---------------------
def X(text, contains=False):
    if contains:
        return f"//*[contains(normalize-space(text()), '{text}')] | //input[contains(@value, '{text}')]"
    return f"//*[normalize-space(text())='{text}'] | //input[@value='{text}']"


def visible(xp):
    out = []
    for e in driver.find_elements(By.XPATH, xp):
        try:
            if e.is_displayed():
                out.append(e)
        except StaleElementReferenceException:
            pass
    return out


def wait_for(xp, timeout=T_WAIT):
    end = time.time() + timeout
    while time.time() < end:
        els = visible(xp)
        if els:
            return els
        time.sleep(0.2)
    raise TimeoutException(f"Not found: {xp}")


def wait_gone(xp, timeout=T_WAIT):
    end = time.time() + timeout
    while time.time() < end:
        if not visible(xp):
            return
        time.sleep(0.2)
    raise TimeoutException(f"Did not close: {xp}")


def click(el):
    for _ in range(6):
        try:
            el.click()
            return
        except (ElementClickInterceptedException, ElementNotInteractableException,
                StaleElementReferenceException):
            time.sleep(0.5)
    driver.execute_script("arguments[0].click();", el)


def click_text(text, idx=-1, contains=False, timeout=T_WAIT):
    els = wait_for(X(text, contains), timeout)
    click(els[idx])
    time.sleep(PAUSE)


def input_after(label, title=None, idx=-1, timeout=T_WAIT):
    """First input after a label; if title is given, search inside that dialog first."""
    cands = []
    if title:
        cands.append(
            f"//*[normalize-space(text())='{title}']"
            f"/following::*[normalize-space(text())='{label}'][1]"
            f"/following::input[not(@type='hidden')][1]"
        )
    cands.append(
        f"//*[normalize-space(text())='{label}']/following::input[not(@type='hidden')][1]"
    )
    end = time.time() + timeout
    while time.time() < end:
        for xp in cands:
            els = visible(xp)
            if els:
                return els[idx]
        time.sleep(0.2)
    raise TimeoutException(f"Input not found after '{label}'")


def set_value(el, text):
    click(el)
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.DELETE)
    el.send_keys(text)
    el.send_keys(Keys.TAB)
    time.sleep(0.2)
    if (el.get_attribute("value") or "").strip() != text:
        # fallback: set the value through JavaScript
        driver.execute_script(
            "arguments[0].value = arguments[1];"
            "arguments[0].dispatchEvent(new Event('input', {bubbles:true}));"
            "arguments[0].dispatchEvent(new Event('change', {bubbles:true}));",
            el, text,
        )


def make_driver():
    if BROWSER.lower() == "chrome":
        d = webdriver.Chrome()
    else:
        d = webdriver.Edge()
    d.maximize_window()
    return d


# ---------------------  Site workflow  ---------------------
def login(user, password):
    driver.get(BASE_URL)
    pwd = wait_for("//input[@type='password']", 30)[0]
    usr = wait_for("//input[@type='text' or not(@type)]", 30)[0]
    usr.clear()
    usr.send_keys(user)
    pwd.clear()
    pwd.send_keys(password)
    btn = visible("//button[normalize-space(.)='Вход'] | //input[@value='Вход'] | //a[normalize-space(.)='Вход']")
    if btn:
        click(btn[-1])
    else:
        pwd.send_keys(Keys.ENTER)
    wait_for(X("Абоненты"), 40)


def recover(user, password):
    """Bring the site back to a clean state after an error."""
    try:
        driver.get(BASE_URL + "/home")
        time.sleep(1.5)
        if visible("//input[@type='password']"):
            login(user, password)
        else:
            wait_for(X("Абоненты"), 30)
    except Exception as e:  # noqa
        print("   (recovery problem:", e, ")")


ACCOUNT_LABEL = "Лицевой счет"


def open_search():
    """Subscribers -> Individuals: open the search form."""
    if not visible(X("Абоненты ФЛ")):
        click_text("Абоненты")
        wait_for(X("Абоненты ФЛ"))
    click_text("Абоненты ФЛ")
    try:
        wait_for(f"//*[normalize-space(text())='{ACCOUNT_LABEL}']", 6)
    except TimeoutException:
        driver.get(SEARCH_URL)
        wait_for(f"//*[normalize-space(text())='{ACCOUNT_LABEL}']", T_WAIT)


def card_residents(timeout=T_WAIT):
    """Current residents value shown on the subscriber card."""
    el = input_after("Проживает", idx=0, timeout=timeout)
    for _ in range(10):
        v = (el.get_attribute("value") or "").strip()
        if v != "":
            return v
        time.sleep(0.5)
    return ""


def process_account(acc, count, today, state):
    """Return (status, message). state["stage"] tracks the current step."""
    state["stage"] = "search"
    open_search()
    inp = input_after(ACCOUNT_LABEL, idx=0)
    set_value(inp, acc)
    # tick the account-number checkbox if it is not ticked
    try:
        cb = visible(f"//*[normalize-space(text())='{ACCOUNT_LABEL}']/preceding::input[@type='checkbox'][1]")
        if cb and not cb[0].is_selected():
            click(cb[0])
        elif not cb:
            click_text(ACCOUNT_LABEL, idx=0)
    except Exception:
        pass
    click_text("Применить")

    state["stage"] = "card"
    try:
        wait_for(X("Рег. док", contains=True), 20)
    except TimeoutException:
        return "not_found", "Subscriber card did not open"

    cur = card_residents()
    if EXPECTED_OLD is not None and cur != EXPECTED_OLD:
        return "skipped_not_1", f"Проживает={cur!r}"

    # ---- Рег. док-та ----
    state["stage"] = "reg_dok"
    click_text("Рег. док", idx=0, contains=True)
    wait_for(X("Рег номер"))
    num = input_after("Рег номер", title=MODAL_REG_TITLE)
    set_value(num, REG_NUMBER)
    rdate = input_after("Рег дата", title=MODAL_REG_TITLE)
    if (rdate.get_attribute("value") or "").strip() != today:
        set_value(rdate, today)
    click_text("Применить")
    click_text("Да")
    wait_gone(X("Рег номер"))
    time.sleep(0.8)

    # ---- Residents count ----
    state["stage"] = "jon_soni"   # an error here means the registration document already exists
    btn_xp = ("//*[normalize-space(text())='Прописано']"
              "/following::*[normalize-space(text())='...' or @value='...'][1]")
    click(wait_for(btn_xp)[0])
    time.sleep(PAUSE)
    wait_for(X(MODAL_RES_TITLE))
    res = input_after("Проживает", title=MODAL_RES_TITLE)
    set_value(res, str(count))
    to = input_after("до", title=MODAL_RES_TITLE)
    if not (to.get_attribute("value") or "").strip():
        set_value(to, today)
    click_text("Применить")
    click_text("Да")
    wait_gone(X(MODAL_RES_TITLE))

    # ---- Verification ----
    state["stage"] = "verify"
    ok = False
    for _ in range(12):
        try:
            if card_residents(timeout=2) == str(count):
                ok = True
                break
        except TimeoutException:
            pass
        time.sleep(0.5)
    if ok:
        return "done", ""
    return "unverified", "Saved, but the new value is not shown on the card - check manually"


def dump_debug(acc, stage):
    ERRORS_DIR.mkdir(exist_ok=True)
    base = ERRORS_DIR / f"{acc}_{stage}_{dt.datetime.now().strftime('%H%M%S')}"
    try:
        driver.save_screenshot(str(base) + ".png")
        Path(str(base) + ".html").write_text(driver.page_source, encoding="utf-8")
    except Exception:
        pass


# ---------------------  Main  ---------------------
def main():
    global driver
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", help="path to the Excel file (default: accounts.xlsx)")
    ap.add_argument("--limit", type=int, help="process only this many accounts (for testing)")
    ap.add_argument("--dry-run", action="store_true", help="no browser: show accounts and planned values")
    ap.add_argument("--all", action="store_true", help="process all accounts (without --all only 3 are processed)")
    args = ap.parse_args()
    if not args.all and not args.limit:
        args.limit = 3
        print("TEST MODE: only 3 accounts. To process everything: python set_residents.py --all")

    excel = Path(args.excel) if args.excel else EXCEL_PATH
    if not excel.exists():
        sys.exit(f"Excel file not found: {excel}")

    accounts = read_accounts(excel)
    print(f"Visible accounts in the workbook: {len(accounts)}")
    if not accounts:
        sys.exit("No accounts found: check ACCOUNT_COL / FIRST_DATA_ROW.")

    if args.dry_run:
        for i, acc in enumerate(accounts[:60]):
            print(f"{i + 1:>4}. {acc} -> {value_for(i)}")
        counts = {}
        for i in range(len(accounts)):
            counts[value_for(i)] = counts.get(value_for(i), 0) + 1
        print("Totals:", dict(sorted(counts.items())))
        return

    cfg = CFG
    user, password = cfg.get("WG_LOGIN"), cfg.get("WG_PASSWORD")
    if not user or not password:
        sys.exit(f"{CRED_FILE.name} must contain WG_LOGIN and WG_PASSWORD (see .env.example).")

    progress = load_progress()
    todo = [(i, a) for i, a in enumerate(accounts) if progress.get(a) not in FINAL_STATUSES]
    if args.limit:
        todo = todo[: args.limit]
    print(f"To process: {len(todo)} (already finished accounts are skipped)")
    if not todo:
        return

    driver = make_driver()
    stats, consecutive_err = {}, 0
    try:
        login(user, password)
        today = dt.date.today().strftime("%d.%m.%Y")
        for n, (pos, acc) in enumerate(todo, 1):
            count = value_for(pos)
            state = {"stage": "?"}
            try:
                st, msg = process_account(acc, count, today, state)
                consecutive_err = 0 if st != "not_found" else consecutive_err
            except Exception as e:  # noqa
                dump_debug(acc, state["stage"])
                # registration document already created -> do not retry (avoids duplicates)
                st = "error_partial" if state["stage"] == "jon_soni" else "error"
                msg = f"[{state['stage']}] {type(e).__name__}: {str(e)[:200]}"
                consecutive_err += 1
                recover(user, password)
            log_progress(acc, pos, st, count, msg)
            stats[st] = stats.get(st, 0) + 1
            print(f"[{n}/{len(todo)}] {acc} -> {count}: {st} {msg}")
            if consecutive_err >= MAX_CONSECUTIVE_ERRORS:
                print("Too many errors in a row. Stopped. See the errors/ folder.")
                break
            time.sleep(PAUSE)
    except KeyboardInterrupt:
        print("\nStopped (Ctrl+C). Run again to resume.")
    finally:
        print("Result:", stats)
        try:
            driver.quit()
        except Exception:
            pass


if __name__ == "__main__":
    main()
